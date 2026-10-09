"""End-to-End: synthetische Mini-Excel -> Validierung -> Cache -> Dashboard-HTML.

Keine echten Daten: die Excel wird im Test erzeugt. Deckt ausserdem die
Randfaelle ab (0 offene Auftraege, Jobs ohne PLZ, Techniker ohne Skills) und
die Eingabevalidierung (fehlendes Sheet/Spalte, kaputte Datei).
"""

from __future__ import annotations

import json

import openpyxl
import pytest

import api.smax_cache as smax_cache
import reporting.dashboard as dashboard
from api.smax_validierung import SmaxImportFehler, validiere_smax_xlsx

OPEN_KOPF = ["Work Order Number", "Account", "City", "Zip", "Model Number",
             "Next PM Due Date", "Order Status", "On Hold Reason"]


def _mini_excel(pfad, *, offene=None, ohne_sheet=None, offen_kopf=None,
                techniker_zeilen=None, offen_account="Klinik A"):
    """Erzeugt eine kleine SMax-Excel. offene = Liste von (id, plz)."""
    offene = [] if offene is None else offene
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("1_Skills Matrix")
    ws.append(["Model Code", "Anna Test", "Bert Test"])
    ws.append(["MC-1", "JA", "NEIN"])

    ws = wb.create_sheet("2_Durchscnittliche_Zeit_MC")
    ws.append(["Zeilenbeschriftungen", "Mittelwert", "Median"])
    ws.append(["MC-1", 1.5, 1.0])

    ws = wb.create_sheet("3_Closed Jobs")
    ws.append(["Work Order Number", "Account", "City", "Zip", "Model Number",
               "Next PM Due Date", "Actual Resolution", "Serial Number",
               "Warranty End Date", "Technician"])
    ws.append(["W1", "Klinik A", "Ulm", "89073", "MC-1", None, "2025-03-01", None, None, "Anna Test"])
    ws.append(["W2", "Klinik A", "Ulm", "89073", "MC-1", None, "2026-03-01", None, None, "Anna Test"])

    ws = wb.create_sheet("4_Open_Jobs")
    ws.append(offen_kopf or OPEN_KOPF)
    for oid, plz in offene:
        ws.append([oid, offen_account, "Ulm", plz, "MC-1", "2026-11-03", "Scheduled", None])

    ws = wb.create_sheet("5_Wohnorte")
    ws.append(["5", "Name", "Land", "Stadt"])
    for zeile in (techniker_zeilen or [[1, "Anna Test", "Germany", "Balingen"],
                                       [2, "Bert Test", "Germany", "Ulm"]]):
        ws.append(zeile)

    if ohne_sheet:
        del wb[ohne_sheet]
    wb.save(pfad)
    return pfad


@pytest.fixture
def isoliert(tmp_path, monkeypatch):
    """Leitet Cache und HTML-Ausgabe ins tmp-Verzeichnis, stellt Dashboard-Globals wieder her."""
    monkeypatch.setattr(smax_cache, "_CACHE", tmp_path / "cache.json")
    monkeypatch.setattr(smax_cache, "_DEMO_MODUS", False)
    monkeypatch.setattr(dashboard, "_OUT_PATH", tmp_path / "tracked_reporting.html")
    monkeypatch.setattr(dashboard, "_OUT_PATH_ROOT", tmp_path / "root.html")
    monkeypatch.setattr(dashboard, "_OUT_PATH_LOKAL", tmp_path / "out.html")
    for name in ("_ECHTDATEN", "_TECH_FARBEN", "_GEBIET_AKTUELL", "_GEBIET_OPTIMIERT",
                 "_NRW_TECHNIKER", "_HUGO_KA_IDS"):
        monkeypatch.setattr(dashboard, name, getattr(dashboard, name))
    return tmp_path


def _pipeline(tmp_path, xlsx):
    data = smax_cache.save_dashboard_data(xlsx)
    dashboard.main([])
    return data, (tmp_path / "out.html").read_text(encoding="utf-8")


# -- End-to-End ---------------------------------------------------------------

def test_e2e_import_cache_dashboard(isoliert):
    xlsx = _mini_excel(isoliert / "t.xlsx", offene=[("OFF-1", "89073"), ("OFF-2", "89073")])
    _, html = _pipeline(isoliert, xlsx)

    cache = json.loads((isoliert / "cache.json").read_text(encoding="utf-8"))
    assert len(cache["techniker"]) == 2
    assert cache["open_jobs"] == 2 and cache["closed_jobs"] == 2
    assert cache["offene_ohne_standort"] == 0
    assert "OFF-1" in html and "OFF-2" in html
    assert "Offene Aufträge (Top 10 nach Fälligkeit)" in html
    assert 'class="planungshinweis"' not in html
    # Echtdaten-Build schreibt NUR nach lokal/, nie in die getrackten (oeffentlichen) Dateien
    assert not (isoliert / "root.html").exists()
    assert not (isoliert / "tracked_reporting.html").exists()


# -- Randfaelle ---------------------------------------------------------------

def test_null_offene_auftraege_zeigt_hinweis(isoliert):
    xlsx = _mini_excel(isoliert / "t.xlsx", offene=[])
    data, html = _pipeline(isoliert, xlsx)
    assert data["open_jobs"] == 0
    assert "Keine offenen Aufträge im aktuellen SMax-Export" in html


def test_jobs_ohne_plz_hinweis_statt_absturz(isoliert):
    xlsx = _mini_excel(isoliert / "t.xlsx",
                       offene=[("OFF-1", "89073"), ("OFF-2", None), ("OFF-3", "01011")])
    data, html = _pipeline(isoliert, xlsx)
    assert data["offene_ohne_standort"] == 2  # leere PLZ + CH-PLZ (Lausanne)
    assert "2 von 3 offenen Aufträgen ohne auflösbaren Standort" in html


def test_techniker_ohne_skills_kein_absturz(isoliert):
    zeilen = [[1, "Anna Test", "Germany", "Balingen"], [2, "Carl Fremd", "Germany", "Essen"]]
    xlsx = _mini_excel(isoliert / "t.xlsx", offene=[("OFF-1", "89073")], techniker_zeilen=zeilen)
    data, html = _pipeline(isoliert, xlsx)
    carl = [t for t in data["techniker"] if t["standort"].lower() == "essen"]
    assert carl and carl[0]["in_skills_matrix"] is False and carl[0]["pm_count"] == 0
    assert "OFF-1" in html


# -- Eingabevalidierung -------------------------------------------------------

def test_validierung_ok(tmp_path):
    xlsx = _mini_excel(tmp_path / "t.xlsx")
    assert validiere_smax_xlsx(xlsx) == []


def test_validierung_fehlendes_sheet_nennt_namen(tmp_path):
    xlsx = _mini_excel(tmp_path / "t.xlsx", ohne_sheet="4_Open_Jobs")
    probleme = validiere_smax_xlsx(xlsx)
    assert any("4_Open_Jobs" in p and "fehlt" in p for p in probleme)


def test_validierung_fehlende_spalte_nennt_sheet_und_spalte(tmp_path):
    kopf = [c for c in OPEN_KOPF if c != "Zip"]
    xlsx = _mini_excel(tmp_path / "t.xlsx", offen_kopf=kopf)
    probleme = validiere_smax_xlsx(xlsx)
    assert any("4_Open_Jobs" in p and "'Zip'" in p for p in probleme)


def test_validierung_leere_arbeitsmappe(tmp_path):
    pfad = tmp_path / "leer.xlsx"
    openpyxl.Workbook().save(pfad)
    assert len(validiere_smax_xlsx(pfad)) >= 5  # alle 5 Sheets fehlen


def test_validierung_keine_xlsx(tmp_path):
    pfad = tmp_path / "kaputt.xlsx"
    pfad.write_text("das ist kein excel", encoding="utf-8")
    probleme = validiere_smax_xlsx(pfad)
    assert probleme and "keine lesbare .xlsx" in probleme[0]


def test_import_bricht_bei_formatfehler_ab_und_laesst_cache_unberuehrt(isoliert):
    (isoliert / "cache.json").write_text('{"alt": true}', encoding="utf-8")
    xlsx = _mini_excel(isoliert / "t.xlsx", ohne_sheet="3_Closed Jobs")
    with pytest.raises(SmaxImportFehler) as exc:
        smax_cache.save_dashboard_data(xlsx)
    assert "3_Closed Jobs" in str(exc.value)
    assert (isoliert / "cache.json").read_text(encoding="utf-8") == '{"alt": true}'


def test_kein_deutscher_techniker_ist_fehler(isoliert):
    xlsx = _mini_excel(isoliert / "t.xlsx", techniker_zeilen=[[1, "Anna Test", "Austria", "Wien"]])
    with pytest.raises(SmaxImportFehler) as exc:
        smax_cache.save_dashboard_data(xlsx)
    assert "Germany" in str(exc.value)


def test_cli_exitcodes(isoliert, capsys):
    assert smax_cache.main([str(isoliert / "gibtsnicht.xlsx")]) == 2
    assert "nicht gefunden" in capsys.readouterr().out
    xlsx = _mini_excel(isoliert / "t.xlsx", ohne_sheet="5_Wohnorte")
    assert smax_cache.main([str(xlsx)]) == 3
    out = capsys.readouterr().out
    assert "5_Wohnorte" in out and "Traceback" not in out
    ok = _mini_excel(isoliert / "ok.xlsx", offene=[("OFF-1", "89073")])
    assert smax_cache.main([str(ok)]) == 0
    assert "ZUSAMMENFASSUNG IMPORT" in capsys.readouterr().out


# -- F02: Ausgabe absichern ---------------------------------------------------

def test_importtext_wird_escaped_in_html_attributen_und_json(isoliert):
    boese_klinik = '<script>alert(1)</script>" onmouseover="alert(2)'
    boese_id = "X'\"><img src=x onerror=alert(3)>"
    xlsx = _mini_excel(isoliert / "t.xlsx", offene=[(boese_id, "89073"), ("OFF-2", "89073")], offen_account=boese_klinik)
    _, html = _pipeline(isoliert, xlsx)

    # kein rohes Tag aus dem Import, keine ausbrechenden Attribute
    assert "<script>alert(1)" not in html
    assert "<img src=x" not in html
    assert '" onmouseover="alert(2)' not in html
    # escaped Darstellung vorhanden (Tabelle)
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in html
    # Inline-Handler enthaelt keine Importdaten
    assert 'onclick="showEmpfehlung(this.dataset.aid)"' in html
    assert "showEmpfehlung('" not in html
    # eingebettetes JSON: kein "</" aus Importdaten
    assert "</script>alert" not in html


def test_js_json_neutralisiert_script_ende():
    out = dashboard._js_json({"k": "</script><!-- <script>"})
    assert "<" not in out
    assert json.loads(out) == {"k": "</script><!-- <script>"}


# -- F03: Claude-Chat im Echtdaten-Modus gesperrt ------------------------------

def test_chat_im_echtdaten_modus_gesperrt(isoliert):
    xlsx = _mini_excel(isoliert / "t.xlsx", offene=[("OFF-1", "89073")])
    _, html = _pipeline(isoliert, xlsx)
    assert 'id="chat-gesperrt-hinweis"' in html
    assert "if (true) { return; }" in html
    assert '<button class="api-key-btn"' not in html
    assert "var SYSTEM_PROMPT = \"\";" in html          # keine Datenzusammenfassung eingebettet
    assert 'id="erklaer-box"' in html                    # Template-Erklaerungen bleiben


def test_chat_freigabe_per_konstante(isoliert, monkeypatch):
    import config
    monkeypatch.setattr(config, "CHAT_MIT_ECHTDATEN_ERLAUBT", True)
    xlsx = _mini_excel(isoliert / "t.xlsx", offene=[("OFF-1", "89073")])
    _, html = _pipeline(isoliert, xlsx)
    assert 'id="chat-gesperrt-hinweis"' not in html
    assert "if (false) { return; }" in html
    assert '<button class="api-key-btn"' in html
