"""Tests: SMax-Header mit 'Objekt: Feld'-Praefix, Datumsparser, Anzeige der Spreizung."""

import datetime
import html
import io
import re

import openpyxl

import reporting.dashboard as dash
from api.import_real_data import (
    baue_zeilen_dict,
    datum_text,
    hole_feld,
    map_closed_job_row,
    map_open_job_row,
    parse_datum_wert,
    parse_smax_xlsx,
    pruefe_header,
    zaehle_unlesbare_datums,
)


class TestHeader:
    def test_header_mit_praefix_wird_unter_beiden_namen_abgelegt(self):
        d = baue_zeilen_dict(["Work Order: Work Order Number", "Account"], ["WO-1", "Acc"])
        assert d["Work Order: Work Order Number"] == "WO-1"
        assert d["Work Order Number"] == "WO-1"
        assert d["Account"] == "Acc"

    def test_header_ohne_praefix(self):
        d = baue_zeilen_dict(["Work Order Number"], ["WO-2"])
        assert hole_feld(d, "Work Order Number") == "WO-2"

    def test_gross_kleinschreibung_tolerant(self):
        d = baue_zeilen_dict(["work order: WORK ORDER NUMBER"], ["WO-3"])
        assert hole_feld(d, "Work Order Number") == "WO-3"

    def test_echte_auftragsnummer_statt_platzhalter(self):
        d = baue_zeilen_dict(
            ["Work Order: Work Order Number", "Account", "City", "Model Number"],
            ["WO-00123", "A", "B", "MC-1"])
        assert map_closed_job_row(d).auftragsnummer == "WO-00123"
        assert map_open_job_row(d).auftragsnummer == "WO-00123"

    def test_leere_nummer_bleibt_leer_fuer_platzhalter(self):
        d = baue_zeilen_dict(["Work Order: Work Order Number"], [None])
        assert map_open_job_row(d).auftragsnummer == ""

    def test_diagnose_nennt_fehlendes_feld_mit_aehnlichstem_header(self):
        w = pruefe_header("4_Open_Jobs", ["Auftrag Nummer", "Account"])
        assert any("Work Order Number" in x and "aehnlichster Header" in x for x in w)

    def test_diagnose_praefix_gilt_als_gefunden(self):
        hdrs = ["Work Order: Work Order Number", "Account", "City", "Zip", "Model Number",
                "Next PM Due Date", "Serial Number", "Warranty End Date", "Contract End Date",
                "Order Status", "On Hold Reason"]
        assert pruefe_header("4_Open_Jobs", hdrs) == []


class TestDatum:
    def test_datetime_zelle(self):
        assert datum_text(datetime.datetime(2026, 10, 5)) == "05.10.2026"
        assert datum_text(datetime.datetime(2026, 10, 5, 14, 30)) == "05.10.2026, 14:30"
        assert datum_text(datetime.date(2026, 10, 5)) == "05.10.2026"

    def test_iso_text(self):
        assert datum_text("2026-10-05") == "05.10.2026"
        assert datum_text("2026-10-05T14:30:00") == "05.10.2026, 14:30"

    def test_deutsches_datum_mit_und_ohne_uhrzeit(self):
        assert datum_text("05.10.2026") == "05.10.2026"
        assert datum_text("05.10.2026, 14:30") == "05.10.2026, 14:30"
        assert parse_datum_wert("05.10.2026 14:30:15") == datetime.datetime(2026, 10, 5, 14, 30, 15)

    def test_leere_zelle(self):
        assert datum_text(None) is None and datum_text("  ") is None
        assert parse_datum_wert("") is None

    def test_unlesbar_wird_gezaehlt(self):
        d = baue_zeilen_dict(["Work Order Number", "Next PM Due Date"], ["W1", "irgendwann"])
        job = map_open_job_row(d)
        assert job.next_pm_due_date == "irgendwann"
        assert zaehle_unlesbare_datums([job], ["next_pm_due_date"]) == {"next_pm_due_date": 1}

    def test_smax_cache_parser_versteht_alle_formate(self):
        from api.smax_cache import _parse_datum
        assert _parse_datum("05.10.2026, 14:30").date() == datetime.date(2026, 10, 5)
        assert _parse_datum("2026-10-05").date() == datetime.date(2026, 10, 5)
        assert _parse_datum(datetime.datetime(2026, 10, 5)).date() == datetime.date(2026, 10, 5)
        assert _parse_datum(None) is None


def _xlsx() -> bytes:
    wb = openpyxl.Workbook()
    wb.active.title = "1_Skills Matrix"
    ws3 = wb.create_sheet("3_Closed Jobs")
    ws3.append(["Work Order: Work Order Number", "Account", "City", "Zip", "Model Number",
                "Next PM Due Date", "Actual Resolution", "Serial Number", "Warranty End Date", "Technician"])
    ws3.append(["WO-1", "A", "Ort", "12345", "MC-1", datetime.datetime(2026, 1, 2),
                "03.01.2026, 10:15", "S1", "2027-01-01", "X Y"])
    ws4 = wb.create_sheet("4_Open_Jobs")
    ws4.append(["Work Order: Work Order Number", "Account", "City", "Zip", "Model Number",
                "Next PM Due Date", "Serial Number", "Warranty End Date", "Contract End Date",
                "Order Status", "On Hold Reason"])
    ws4.append(["WO-2", "A", "Ort", "12345", "MC-1", None, "S1", None, None, "Scheduled", None])
    ws4.append(["WO-3", "A", "Ort", "12345", "MC-1", "10.02.2026", "S2", None, None, "Scheduled", None])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestImportEndToEnd:
    def test_echte_nummern_und_faelligkeit(self):
        erg = parse_smax_xlsx(_xlsx(), sample_limit=None)
        assert [a.auftragsnummer for a in erg.offene_auftraege] == ["WO-2", "WO-3"]
        assert erg.geschlossene_auftraege[0].auftragsnummer == "WO-1"
        assert erg.geschlossene_auftraege[0].next_pm_due_date == "02.01.2026"
        assert erg.geschlossene_auftraege[0].erledigung_datum == "03.01.2026, 10:15"
        assert sum(1 for a in erg.offene_auftraege if a.next_pm_due_date) == 1

    def test_leere_faelligkeit_wird_gemeldet(self):
        erg = parse_smax_xlsx(_xlsx(), sample_limit=None)
        warn = " ".join(w for sr in erg.sheet_ergebnisse for w in sr.warnungen)
        assert "1 von 2 offenen Auftraegen ohne Faelligkeit" in warn
        assert "Header-Feld" not in warn


def _metrik(tid, kliniken, pct):
    return {"id": tid, "standort": "Ort", "kliniken": kliniken, "avg_fahrzeit": 30 if kliniken else 0,
            "max_fahrzeit": 40, "fahrtstunden_jahr": 10, "onsite_stunden": 10, "ratio": 1.0,
            "auslastung_pct": pct, "plz_info": None, "alternative": None,
            "verschoben": 0, "verschoben_gewonnen": 0, "verschoben_abgegeben": 0, "abgaben": []}


def _kopf(seite: str) -> str:
    kopf = seite.split('id="go-view-optimiert"')[1].split("</p>")[0]
    return html.unescape(re.sub(r"<[^>]+>", "", kopf))


class TestSpreizungOhneTechnikerOhneSkills:
    def test_zwei_spreizungen_und_nicht_beruecksichtigt(self):
        akt = [_metrik("A", 10, 40), _metrik("B", 10, 20), _metrik("C", 0, 0), _metrik("D", 0, 0)]
        opt = [_metrik("A", 10, 35), _metrik("B", 10, 25), _metrik("C", 0, 0), _metrik("D", 0, 0)]
        techniker = {t: {"standort": "Ort", "lat": 50.0, "lon": 9.0, "in_skills_matrix": t in "AB"}
                     for t in "ABCD"}
        text = _kopf(dash._render_gebietsoptimierung(akt, opt, techniker, gebiete_status={}))
        assert "alle Techniker: Median vorher" in text
        assert "Spreizung vorher 40" in text               # inkl. Techniker ohne Klinik
        assert "Nur Techniker mit Klinik: Median vorher 30" in text
        assert "Spreizung vorher 20, nachher 10" in text   # nur A und B
        assert "Nicht berücksichtigt: 2 Techniker ohne Klinik (2 wegen fehlender Skills)" in text

    def test_hinweis_zum_median_uebersetzbar(self):
        akt = [_metrik("A", 10, 40)]
        seite = dash._render_gebietsoptimierung(
            akt, akt, {"A": {"standort": "O", "lat": 1.0, "lon": 1.0}}, gebiete_status={})
        assert "Der Median kann sinken" in seite
        for key in ("Nicht berücksichtigt:", "wegen fehlender Skills", "Techniker ohne Klinik",
                    "Nur Techniker mit Klinik: Median vorher"):
            assert key in dash.LABEL_MAP_EN
        assert any("Der Median kann sinken" in k for k in dash.LABEL_MAP_EN)

    def test_ohne_techniker_ohne_klinik_nur_ein_wertepaar(self):
        akt = [_metrik("A", 10, 40), _metrik("B", 10, 20)]
        techniker = {t: {"standort": "O", "lat": 1.0, "lon": 1.0} for t in "AB"}
        text = _kopf(dash._render_gebietsoptimierung(akt, akt, techniker, gebiete_status={}))
        assert "Nicht berücksichtigt" not in text
