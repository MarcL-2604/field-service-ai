"""
api/smax_validierung.py
=======================
Eingabevalidierung fuer den SMax-Excel-Export (KI_gestuetzte_Planung.xlsx).

Prueft VOR der Verarbeitung, dass alle 5 Sheets und die Pflichtspalten vorhanden
sind. Statt stiller Falschdaten (leere Sheets -> stiller Demo-Fallback im
Dashboard) liefert der Import so eine verstaendliche deutsche Fehlermeldung mit
konkretem Sheet-/Spaltennamen.

Oeffentliche API:
  SmaxImportFehler             -- Exception mit deutscher, traceback-freier Meldung
  validiere_smax_xlsx(quelle)  -- Liste aller gefundenen Probleme (leer = ok)
  pruefe_smax_xlsx(quelle)     -- wirft SmaxImportFehler bei Problemen
"""

from __future__ import annotations

import io
from pathlib import Path
from typing import Union

SHEET_SKILLS = "1_Skills Matrix"
SHEET_ZEITEN = "2_Durchscnittliche_Zeit_MC"  # Schreibweise ohne "h" ist im Export so
SHEET_CLOSED = "3_Closed Jobs"
SHEET_OPEN = "4_Open_Jobs"
SHEET_WOHNORTE = "5_Wohnorte"

# Pflichtspalten je Sheet. Eine Alternative (Tupel) ist erfuellt, wenn EINER der
# Namen vorhanden ist (z.B. Pivot-Artefakt "Zeilenbeschriftungen").
_PFLICHT_SPALTEN: dict[str, list[tuple[str, ...]]] = {
    SHEET_ZEITEN: [("Model_Code", "Zeilenbeschriftungen"), ("Mittelwert",), ("Median",)],
    SHEET_CLOSED: [("Account",), ("City",), ("Zip",), ("Model Number",),
                   ("Actual Resolution",), ("Technician",)],
    SHEET_OPEN: [("Account",), ("City",), ("Zip",), ("Model Number",),
                 ("Next PM Due Date",), ("Order Status",)],
    SHEET_WOHNORTE: [("Name",), ("Land",), ("Stadt",)],
}

# Sheets, die mindestens eine Datenzeile enthalten muessen. 4_Open_Jobs darf
# leer sein (0 offene Auftraege ist ein gueltiger Zustand).
_BRAUCHT_DATEN = (SHEET_SKILLS, SHEET_CLOSED, SHEET_WOHNORTE)

_MIN_TECHNIKER_SPALTEN = 1

Quelle = Union[str, Path, bytes]


class SmaxImportFehler(Exception):
    """Fehler beim SMax-Import; str(e) ist eine fertige deutsche Meldung."""

    def __init__(self, probleme: list[str]) -> None:
        self.probleme = probleme
        super().__init__("\n".join(f"  - {p}" for p in probleme))


def _oeffnen(quelle: Quelle):
    import openpyxl
    if isinstance(quelle, (bytes, bytearray)):
        return openpyxl.load_workbook(io.BytesIO(bytes(quelle)), read_only=True, data_only=True)
    return openpyxl.load_workbook(str(quelle), read_only=True, data_only=True)


def validiere_smax_xlsx(quelle: Quelle) -> list[str]:
    """Gibt alle gefundenen Probleme als deutsche Meldungen zurueck (leer = ok)."""
    try:
        wb = _oeffnen(quelle)
    except FileNotFoundError:
        return [f"Datei nicht gefunden: {quelle}"]
    except Exception as e:  # kaputte/falsche Datei (z.B. .xls, CSV, passwortgeschuetzt)
        return [f"Datei ist keine lesbare .xlsx-Datei ({type(e).__name__}). "
                f"Bitte den SMax-Export als Excel-Arbeitsmappe (.xlsx) speichern."]

    probleme: list[str] = []
    try:
        erwartet = [SHEET_SKILLS, SHEET_ZEITEN, SHEET_CLOSED, SHEET_OPEN, SHEET_WOHNORTE]
        for name in erwartet:
            if name not in wb.sheetnames:
                probleme.append(
                    f"Sheet '{name}' fehlt. Vorhandene Sheets: "
                    f"{', '.join(repr(s) for s in wb.sheetnames) or '(keine)'}. "
                    f"Sheet-Namen muessen exakt stimmen (inkl. Schreibweise).")
                continue

            zeilen = wb[name].iter_rows(values_only=True)
            kopf = next(zeilen, None)
            if kopf is None or not any(c not in (None, "") for c in kopf):
                probleme.append(f"Sheet '{name}' ist leer (keine Kopfzeile).")
                continue
            header = [str(h).strip() for h in kopf if h not in (None, "")]

            if name == SHEET_SKILLS:
                if len(header) - 1 < _MIN_TECHNIKER_SPALTEN:
                    probleme.append(
                        f"Sheet '{name}': keine Techniker-Spalten gefunden "
                        f"(Spalte A = Model Code, ab Spalte B ein Techniker je Spalte).")
            else:
                for alternativen in _PFLICHT_SPALTEN.get(name, []):
                    if not any(a in header for a in alternativen):
                        probleme.append(
                            f"Sheet '{name}': Pflichtspalte '{' / '.join(alternativen)}' fehlt. "
                            f"Gefundene Spalten: {', '.join(header)}")

            if name in _BRAUCHT_DATEN:
                hat_daten = any(
                    any(c not in (None, "") for c in zeile) for zeile in zeilen
                )
                if not hat_daten:
                    probleme.append(f"Sheet '{name}' enthaelt keine Datenzeilen.")
    finally:
        wb.close()
    return probleme


def pruefe_smax_xlsx(quelle: Quelle) -> None:
    """Wirft SmaxImportFehler, wenn die Datei nicht dem erwarteten Format entspricht."""
    probleme = validiere_smax_xlsx(quelle)
    if probleme:
        raise SmaxImportFehler(probleme)
